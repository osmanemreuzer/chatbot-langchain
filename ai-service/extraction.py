from pathlib import Path

from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document


class SafeTextLoader(BaseLoader):
    # LangChain'in hazır TextLoader'ı geçersiz bir UTF-8 baytında direkt
    # UnicodeDecodeError fırlatıyor (ya da encoding tahminine düşüyor, mojibake
    # riski taşıyor). Burada sadece geçersiz baytları "replace" ediyoruz, tüm
    # dosyayı yanlış encoding ile yeniden yorumlamıyoruz.
    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> list[Document]:
        raw = Path(self.file_path).read_bytes()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("utf-8", errors="replace")
        return [Document(page_content=text, metadata={"source": self.file_path})]


class DocxTableAwareLoader(BaseLoader):
    # LangChain'in hazır Docx2txtLoader'ı docx2txt kütüphanesini kullanıyor ve
    # tablo hücrelerini paragraflardan ayrı, satır satır çıkarmıyor. Buradaki
    # python-docx tabanlı mantık öncekiyle birebir aynı (tablo hücreleri " | "
    # ile birleştirilip ayrı satır olarak ekleniyor).
    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> list[Document]:
        document = DocxDocument(self.file_path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [Document(page_content="\n".join(parts), metadata={"source": self.file_path})]


def load_document(file_path: str) -> list[Document]:
    extension = Path(file_path).suffix.lower()

    if extension in (".txt", ".md"):
        return SafeTextLoader(file_path).load()

    if extension == ".pdf":
        return PyPDFLoader(file_path).load()  # sayfa başına bir Document, page metadata dahil

    if extension == ".docx":
        return DocxTableAwareLoader(file_path).load()

    raise ValueError(f"Desteklenmeyen dosya türü: {extension}")
#burada artık pdf.md.txt dosyaları hepsı documentlar seklınde gelıyor.